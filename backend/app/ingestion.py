import csv
import io
import json
import posixpath
import stat
import subprocess
import sys
import zipfile
from dataclasses import dataclass
from pathlib import Path, PurePosixPath

from docx import Document as DocxDocument

from app.security import SECRET_PATTERN, redact
from app.settings import Settings

SUPPORTED = frozenset({".txt", ".md", ".pdf", ".docx", ".json", ".csv"})


class IngestionError(ValueError):
    pass


@dataclass(frozen=True)
class ParsedDocument:
    name: str
    text: str


def _zip_members(data: bytes, settings: Settings) -> list[tuple[str, bytes]]:
    limit = settings.max_decompressed_mb * 1024 * 1024
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            members = archive.infolist()
            if len(members) > settings.max_archive_files:
                raise IngestionError("El archivo ZIP excede el límite de entradas.")
            if sum(member.file_size for member in members) > limit:
                raise IngestionError("El archivo ZIP excede el tamaño descomprimido permitido.")
            output: list[tuple[str, bytes]] = []
            read_bytes = 0
            for member in members:
                path = PurePosixPath(member.filename.replace("\\", "/"))
                if (
                    path.is_absolute()
                    or ".." in path.parts
                    or not path.parts
                    or ":" in path.parts[0]
                    or "\x00" in member.filename
                    or stat.S_ISLNK(member.external_attr >> 16)
                ):
                    raise IngestionError("El archivo ZIP contiene una ruta insegura.")
                if member.flag_bits & 1:
                    raise IngestionError("Los archivos ZIP cifrados no están admitidos.")
                if member.is_dir():
                    continue
                if member.file_size > max(1, member.compress_size) * settings.max_zip_ratio:
                    raise IngestionError("El archivo ZIP excede el límite de compresión.")
                with archive.open(member) as stream:
                    contents = stream.read(limit - read_bytes + 1)
                read_bytes += len(contents)
                if read_bytes > limit:
                    raise IngestionError("El archivo ZIP excede el tamaño descomprimido permitido.")
                output.append((str(path), contents))
            return output
    except (zipfile.BadZipFile, RuntimeError, OSError, NotImplementedError) as exc:
        raise IngestionError("No se pudo leer el archivo ZIP.") from exc


def _decode(data: bytes) -> str:
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1252")


def _parse(name: str, data: bytes, settings: Settings) -> str:
    suffix = PurePosixPath(name).suffix.lower()
    if suffix in {".txt", ".md"}:
        return _decode(data)
    if suffix == ".json":
        return json.dumps(json.loads(_decode(data)), ensure_ascii=False, indent=2)
    if suffix == ".csv":
        return "\n".join(" | ".join(row) for row in csv.reader(io.StringIO(_decode(data))))
    if suffix == ".pdf":
        return _parse_pdf(data, settings)
    if suffix == ".docx":
        _zip_members(data, settings)  # DOCX is ZIP: check decompression before parsing XML.
        document = DocxDocument(io.BytesIO(data))
        parts = [paragraph.text for paragraph in document.paragraphs]
        parts.extend(
            " | ".join(cell.text for cell in row.cells)
            for table in document.tables
            for row in table.rows
        )
        return "\n".join(parts)
    raise IngestionError("Formato no admitido.")


def _parse_pdf(data: bytes, settings: Settings) -> str:
    if not data.startswith(b"%PDF-"):
        raise IngestionError("PDF inválido.")
    limit = settings.max_decompressed_mb * 1024 * 1024
    try:
        result = subprocess.run(
            [
                sys.executable,
                "-I",
                str(Path(__file__).with_name("pdf_parser.py")),
                str(limit),
                str(settings.max_upload_mb * 1024 * 1024),
            ],
            input=data,
            capture_output=True,
            timeout=6,
            check=False,
            start_new_session=True,
            env={"PATH": "/usr/bin:/bin", "LANG": "C.UTF-8"},
        )
    except (OSError, subprocess.TimeoutExpired) as exc:
        raise IngestionError("El lector PDF excedió su límite o no está disponible.") from exc
    if result.returncode != 0 or len(result.stdout) > limit * 4:
        raise IngestionError("PDF inválido o fuera del límite de procesamiento.")
    return result.stdout.decode("utf-8", errors="strict")


def parse_upload(
    name: str, data: bytes, settings: Settings
) -> tuple[list[ParsedDocument], list[str]]:
    # Never create a path or extract an archive. Only safe display names are retained.
    name = redact(posixpath.basename(name.replace("\\", "/"))[:200])
    if not name or len(data) > settings.max_upload_mb * 1024 * 1024:
        raise IngestionError("El archivo excede el tamaño permitido o no tiene nombre.")
    zipped = PurePosixPath(name).suffix.lower() == ".zip"
    inputs = _zip_members(data, settings) if zipped else [(name, data)]
    documents: list[ParsedDocument] = []
    skipped: list[str] = []
    total_characters = 0
    for entry_name, contents in inputs:
        entry_name = redact(entry_name)
        if PurePosixPath(entry_name).suffix.lower() not in SUPPORTED:
            if not zipped:
                raise IngestionError("Formato no admitido: TXT, MD, PDF, DOCX, CSV, JSON o ZIP.")
            skipped.append(f"{entry_name}: formato no admitido")
            continue
        try:
            text = _parse(entry_name, contents, settings).replace("\x00", "").strip()
        except Exception as exc:
            # Parser exceptions can include uploaded content or filesystem paths.
            if not zipped:
                raise IngestionError("No se pudo leer el documento. Verificá su formato.") from exc
            skipped.append(f"{entry_name}: no se pudo leer")
            continue
        total_characters += len(text)
        if total_characters > settings.max_decompressed_mb * 1024 * 1024:
            raise IngestionError("El texto extraído excede el tamaño permitido.")
        if SECRET_PATTERN.search(text):
            if not zipped:
                raise IngestionError(
                    "Se detectaron credenciales. Quitalas antes de cargar el archivo."
                )
            skipped.append(f"{entry_name}: contiene credenciales")
            continue
        if not text:
            if not zipped:
                raise IngestionError("El documento no tiene texto extraíble; no se aplica OCR.")
            skipped.append(f"{entry_name}: sin texto extraíble")
            continue
        documents.append(ParsedDocument(entry_name, text))
    if not documents:
        raise IngestionError("No hay documentos con texto válido en la carga.")
    return documents, skipped


def chunks(text: str, size: int = 1000, overlap: int = 150) -> list[str]:
    if overlap < 0 or size <= overlap:
        raise ValueError("Chunk size must exceed overlap")
    result: list[str] = []
    start = 0
    while start < len(text):
        end = min(start + size, len(text))
        if end < len(text):
            boundary = text.rfind(" ", start + size // 2, end)
            if boundary > start:
                end = boundary
        fragment = text[start:end].strip()
        if fragment:
            result.append(fragment)
        if end == len(text):
            break
        start = max(start + 1, end - overlap)
    return result

import io
import subprocess
import zipfile
from typing import Any

import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import ArrayObject, DecodedStreamObject, DictionaryObject, NameObject

from app.core.settings import Settings
from app.knowledge.ingestion import IngestionError, chunks, parse_upload


def archive(entries: dict[str, bytes], compression: int = zipfile.ZIP_STORED) -> bytes:
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", compression=compression) as zipped:
        for name, contents in entries.items():
            zipped.writestr(name, contents)
    return output.getvalue()


@pytest.mark.parametrize(
    ("name", "contents", "expected"),
    [
        ("notes.txt", b"Horario: lunes a viernes", "lunes"),
        ("notes.md", b"# Planes\nStarter 29 euros", "Starter"),
        ("notes.json", b'{"plan": "Starter", "price": 29}', "Starter"),
        ("notes.csv", b"plan,price\nStarter,29", "Starter | 29"),
        ("legacy.txt", "Espa\xf1a".encode("cp1252"), "España"),
    ],
)
def test_formats(settings: Settings, name: str, contents: bytes, expected: str) -> None:
    documents, skipped = parse_upload(name, contents, settings)
    assert expected in documents[0].text
    assert skipped == []


def test_real_docx_with_tables(settings: Settings) -> None:
    document = Document()
    document.add_paragraph("Soporte Madrid")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Starter", "29 euros"
    output = io.BytesIO()
    document.save(output)
    parsed, _ = parse_upload("rates.docx", output.getvalue(), settings)
    assert "Soporte Madrid" in parsed[0].text
    assert "Starter | 29 euros" in parsed[0].text


def test_real_pdf_text(settings: Settings) -> None:
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=200)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
    )
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 10 100 Td (Forma Starter 29 euros) Tj ET")
    page[NameObject("/Contents")] = ArrayObject([stream])
    output = io.BytesIO()
    writer.write(output)
    parsed, _ = parse_upload("pricing.pdf", output.getvalue(), settings)
    assert "Forma Starter 29 euros" in parsed[0].text


@pytest.mark.parametrize("nested_array", [False, True])
def test_pdf_decompressed_stream_budget(settings: Settings, nested_array: bool) -> None:
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=200)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
    )
    stream = DecodedStreamObject()
    stream.set_data(b" " * (2 * 1024 * 1024) + b"BT /F1 12 Tf 10 100 Td (Company knowledge) Tj ET")
    compressed = stream.flate_encode()
    page[NameObject("/Contents")] = ArrayObject([compressed]) if nested_array else compressed
    output = io.BytesIO()
    writer.write(output)
    assert len(output.getvalue()) < 10000
    allowed, _ = parse_upload("compressed.pdf", output.getvalue(), settings)
    assert allowed[0].text == "Company knowledge"
    small = settings.model_copy(update={"max_decompressed_mb": 1})
    with pytest.raises(IngestionError):
        parse_upload("compressed.pdf", output.getvalue(), small)


def test_pdf_parser_timeout_sanitized(settings: Settings, monkeypatch: pytest.MonkeyPatch) -> None:
    def timeout(*args: Any, **kwargs: Any) -> subprocess.CompletedProcess[bytes]:
        assert kwargs["timeout"] == 6
        assert "ANTHROPIC_API_KEY" not in kwargs["env"]
        raise subprocess.TimeoutExpired(args[0], 6, stderr=b"sensitive debug")

    monkeypatch.setattr(subprocess, "run", timeout)
    with pytest.raises(IngestionError) as failure:
        parse_upload("slow.pdf", b"%PDF-1.7\nplaceholder", settings)
    assert "sensitive" not in str(failure.value)


def test_zip_memory_ingestion_and_skips(settings: Settings) -> None:
    data = archive({"folder/doc.md": b"Knowledge", "nested.zip": b"ignored", "bad.json": b"{"})
    documents, skipped = parse_upload("docs.zip", data, settings)
    assert [document.name for document in documents] == ["folder/doc.md"]
    assert len(skipped) == 2
    assert not (settings.data_dir / "folder").exists()


@pytest.mark.parametrize(
    "path", ["../escape.txt", "/absolute.txt", "a/../../b.txt", "..\\escape.txt", "C:\\secret.txt"]
)
def test_zip_rejects_unsafe_paths(settings: Settings, path: str) -> None:
    with pytest.raises(IngestionError, match="insegura"):
        parse_upload("bad.zip", archive({path: b"text"}), settings)


def test_zip_bomb_and_entry_limit(settings: Settings) -> None:
    with pytest.raises(IngestionError, match="compresión"):
        parse_upload(
            "bomb.zip", archive({"doc.txt": b"x" * 100000}, zipfile.ZIP_DEFLATED), settings
        )
    small = settings.model_copy(update={"max_archive_files": 1})
    with pytest.raises(IngestionError, match="entradas"):
        parse_upload("many.zip", archive({"a.txt": b"a", "b.txt": b"b"}), small)
    small = settings.model_copy(update={"max_decompressed_mb": 1})
    with pytest.raises(IngestionError, match="descomprimido"):
        parse_upload("large.zip", archive({"a.txt": b"a" * (1024 * 1024 + 1)}), small)


def test_zip_symlink(settings: Settings) -> None:
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as zipped:
        member = zipfile.ZipInfo("linked.txt")
        member.create_system = 3
        member.external_attr = 0o120777 << 16
        zipped.writestr(member, b"target")
    with pytest.raises(IngestionError, match="insegura"):
        parse_upload("links.zip", output.getvalue(), settings)


@pytest.mark.parametrize(
    ("name", "data"),
    [
        ("empty.txt", b" "),
        ("script.py", b"print(1)"),
        ("bad.json", b"{"),
        ("broken.pdf", b"not a pdf"),
        ("broken.docx", b"not a docx"),
    ],
)
def test_invalid_upload(settings: Settings, name: str, data: bytes) -> None:
    with pytest.raises(IngestionError):
        parse_upload(name, data, settings)


def test_secret_not_persisted(settings: Settings) -> None:
    with pytest.raises(IngestionError, match="credenciales"):
        parse_upload("bad.txt", b"ANTHROPIC_API_KEY=sk-ant-test-abcdefghijklmnop", settings)
    data = archive({"ok.txt": b"Normal business facts", "key.txt": b"PASSWORD=abcdefghi"})
    documents, skipped = parse_upload("mixed.zip", data, settings)
    assert len(documents) == 1
    assert "credenciales" in skipped[0]


@pytest.mark.parametrize(
    "contents", [b'{"PASSWORD":"exam-placeholder"}', b'{"ANTHROPIC_API_KEY": "exam-placeholder"}']
)
def test_json_secret_field_rejected(settings: Settings, contents: bytes) -> None:
    with pytest.raises(IngestionError, match="credenciales"):
        parse_upload("settings.json", contents, settings)


def test_chunk_overlap_and_end() -> None:
    text = " ".join(f"word{number}" for number in range(600))
    fragments = chunks(text, size=300, overlap=40)
    assert all(len(fragment) <= 300 for fragment in fragments)
    assert fragments[0].startswith("word0")
    assert fragments[-1].endswith("word599")
    assert len(fragments) > 5
    with pytest.raises(ValueError):
        chunks(text, size=20, overlap=20)
